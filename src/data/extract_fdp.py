"""Extract DILG Full Disclosure Policy filings for Mapandan into fdp_disclosures.json.

Reads datasets/fdp/<year>/*.xlsx (FDP Forms 6/7/8/9/10/12/13/14, SRE, SEF,
LBP budget, SIPB indebtedness), normalizes them into quarterly records, and
writes src/data/fdp_disclosures.json for the site generators.

Usage: python3 src/data/extract_fdp.py
"""
import hashlib
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

from src.data._utils import write_json

CACHE_DIR = Path("datasets/fdp")
OUTPUT = Path("src/data/fdp_disclosures.json")

CERTIFY_RE = re.compile(r"we hereby certify", re.I)
CODE_RE = re.compile(r"^\d-\d{2}-\d{2}-\d{3}$")


def to_float(value):
    """Coerce FDP money cells (floats, comma text, 'P 1,234', N/A) to float/None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("\u00b1", "")
    text = re.sub(r"^(p|php|₱)\s*", "", text, flags=re.I).strip()
    if text in ("", "-", "N/A", "n/a", "NONE", "none", "not yet started"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def sheet_rows(ws, max_col=12):
    for row in ws.iter_rows(values_only=True):
        yield [c for c in row[:max_col]]


def find_row(rows, *needles):
    """Return first row whose col-A text starts with any needle (case-insensitive)."""
    for row in rows:
        label = clean_text(row[0]).lower() if len(row) > 0 else ""
        for needle in needles:
            if label.startswith(needle.lower()):
                return row
    return None


def parse_year_quarter(ws, rows):
    """FDP header blocks carry CALENDAR YEAR / QUARTER cells or 'Period Covered:'."""
    year, quarter = None, None
    for row in rows[:8]:
        cells = [clean_text(c) for c in row]
        joined = " ".join(cells)
        if "period covered" in joined.lower():
            match = re.search(r"Q([1-4])\s*,?\s*(20\d{2})", joined, re.I)
            if match:
                quarter, year = int(match.group(1)), int(match.group(2))
        for i, cell in enumerate(cells):
            rest = [c for c in cells[i + 1:] if c]
            if cell.lower() == "calendar year:" and rest:
                try:
                    year = int(float(rest[0]))
                except ValueError:
                    pass
            if cell.lower() == "quarter:" and rest:
                try:
                    quarter = int(float(rest[0]))
                except ValueError:
                    pass
    return year, quarter


def period_key(year, quarter):
    if year and quarter:
        return f"{year}-Q{quarter}"
    if year:
        return str(year)
    return "undated"


def money_row(row, cols):
    """Pull floats from given 0-indexed columns."""
    return [to_float(row[i]) if i < len(row) else None for i in cols]


# ----------------------------------------------------------------------------
# Form parsers (each returns a record dict or None)
# ----------------------------------------------------------------------------

def parse_sre(ws):
    rows = list(sheet_rows(ws))
    year, quarter = parse_year_quarter(ws, rows)
    header_idx = next((i for i, r in enumerate(rows)
                       if clean_text(r[0]).lower().startswith("particulars")
                       or (len(r) > 1 and clean_text(r[1]).lower().startswith("particulars"))), None)
    if header_idx is None:
        return None
    p = 0 if clean_text(rows[header_idx][0]).lower().startswith("particulars") else 1
    records = []
    for row in rows[header_idx + 1:]:
        label = clean_text(row[p]) if len(row) > p else ""
        if not label:
            continue
        nums = [to_float(c) for c in row[p + 1:]]
        vals = [v for v in nums if v is not None]
        if not vals:
            continue
        while len(vals) < 5:
            vals.append(None)
        records.append({"label": label, "target": vals[0], "general_fund": vals[1],
                        "sef": vals[2], "trust_fund": vals[3], "trust_liability": vals[4]})
    if not records:
        return None
    return {"form": "sre", "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "rows": records}


def last_number(row):
    nums = [to_float(c) for c in row]
    for v in reversed(nums):
        if v is not None:
            return v
    return None


def parse_sef(ws):
    rows = list(sheet_rows(ws, 9))
    year, quarter = parse_year_quarter(ws, rows)
    receipt, items, subtotal, balance = None, [], None, None
    in_items = False
    for row in rows:
        a = clean_text(row[0]) if len(row) > 0 else ""
        low = a.lower()
        if low.startswith("receipt from sef"):
            receipt = last_number(row)
            in_items = True
            continue
        if low.startswith("sub"):
            subtotal = last_number(row)
            in_items = False
            continue
        if low.startswith("balance"):
            balance = last_number(row)
            in_items = False
            continue
        if CERTIFY_RE.search(a) or "hereby certify" in " ".join(clean_text(c) for c in row).lower():
            in_items = False
            continue
        if in_items:
            b = clean_text(row[1]) if len(row) > 1 else ""
            amount = last_number(row)
            if b and amount is not None:
                items.append({"label": b, "amount": amount})
    return {"form": "sef", "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "receipt": receipt, "items": items,
            "subtotal": subtotal, "balance": balance}


def parse_ldrrmf(ws):
    rows = list(sheet_rows(ws, 10))
    year, quarter = parse_year_quarter(ws, rows)
    sources, utilization = [], []
    section = None
    for row in rows:
        a = clean_text(row[0]) if len(row) > 0 else ""
        low = a.lower()
        if low.startswith("a. sources of funds"):
            section = "sources"
            continue
        if low.startswith("b. utilization"):
            section = "util"
            continue
        if CERTIFY_RE.search(a):
            break
        if section == "sources" and a:
            qrf, mit = money_row(row, [1, 2])
            total = to_float(row[6]) if len(row) > 6 else None
            if qrf is not None or mit is not None or total is not None:
                sources.append({"label": a, "qrf": qrf, "mitigation": mit, "total": total})
        elif section == "util" and a and not low.startswith("total utilization") and not low.startswith("unutilized"):
            qrf = to_float(row[1]) if len(row) > 1 else None
            total = to_float(row[6]) if len(row) > 6 else None
            if qrf not in (None, 0) or (total not in (None, 0)):
                utilization.append({"label": a, "qrf": qrf, "total": total})
    totals = {}
    for row in rows:
        a = clean_text(row[0]).lower() if len(row) > 0 else ""
        if a.startswith("total utilization"):
            totals["utilization"] = to_float(row[6]) if len(row) > 6 else None
        elif a.startswith("unutilized balance"):
            totals["unutilized"] = to_float(row[6]) if len(row) > 6 else None
        elif a.startswith("total funds avail"):
            totals["available"] = to_float(row[6]) if len(row) > 6 else None
    if not sources and not utilization:
        return None
    return {"form": "ldrrmf", "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "sources": sources, "utilization": utilization,
            "totals": totals}


def parse_scf(ws):
    rows = list(sheet_rows(ws, 13))
    year, quarter = parse_year_quarter(ws, rows)
    section, entries, key = None, [], {}
    for row in rows:
        cells = [clean_text(c) for c in row]
        a = cells[0] if cells else ""
        if a.lower().startswith("cash flows from"):
            section = re.sub(r"cash flows from\s*", "", a, flags=re.I).strip() or section
            continue
        if CERTIFY_RE.search(a):
            break
        label = next((c for c in cells[:4] if c), "")
        if not label:
            continue
        nums = [to_float(c) for c in cells[4:]]
        value = next((v for v in reversed(nums) if v is not None), None)
        low = label.lower()
        entries.append({"label": label, "section": section, "value": value})
        if low.startswith("net cash from operating"):
            key["net_operating"] = value
        elif low.startswith("net cash from investing"):
            key["net_investing"] = value
        elif low.startswith("net cash from financ"):
            key["net_financing"] = value
        elif low.startswith("net increase in cash"):
            key["net_increase"] = value
        elif low.startswith("cash at beginning"):
            key["beginning"] = value
        elif low.startswith("cash at the end"):
            key["ending"] = value
    if not entries:
        return None
    return {"form": "cash_flows", "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "key": key, "rows": entries}


def parse_uca(ws):
    rows = list(sheet_rows(ws, 11))
    year, quarter = parse_year_quarter(ws, rows)
    debtors, total, started = [], None, False
    for row in rows:
        a = clean_text(row[0]) if len(row) > 0 else ""
        if a.lower().startswith("name of debtor"):
            started = True
            continue
        if not started:
            continue
        if CERTIFY_RE.search(a):
            break
        if a.lower().startswith("total"):
            total = to_float(row[1]) if len(row) > 1 else None
            continue
        if not a or a.upper() in ("NONE", "N/A"):
            continue
        debtors.append({"debtor": a,
                        "balance": to_float(row[1]) if len(row) > 1 else None,
                        "date_granted": clean_text(row[2]) if len(row) > 2 else "",
                        "purpose": clean_text(row[3]) if len(row) > 3 else ""})
    if not debtors and total in (None, 0):
        return {"form": "cash_advances", "period": period_key(year, quarter),
                "year": year, "quarter": quarter, "status": "nil",
                "debtors": [], "total": 0}
    return {"form": "cash_advances", "period": period_key(year, quarter),
            "year": year, "quarter": quarter, "debtors": debtors, "total": total}


def parse_trust(ws, suffix):
    rows = list(sheet_rows(ws, 13))
    year, quarter = parse_year_quarter(ws, rows)
    header_idx = next((i for i, r in enumerate(rows)
                       if clean_text(r[0]).lower().startswith(("program or project", "fund source"))), None)
    items = []
    if header_idx is not None:
        for row in rows[header_idx + 2:]:
            a = clean_text(row[0]) if len(row) > 0 else ""
            if CERTIFY_RE.search(" ".join(clean_text(c) for c in row)) or a.lower().startswith("certified correct"):
                break
            if not a or a.upper() in ("N/A", "NONE"):
                continue
            items.append({"program": a,
                          "location": clean_text(row[1]) if len(row) > 1 else "",
                          "cost": to_float(row[2]) if len(row) > 2 else None,
                          "started": clean_text(row[3]) if len(row) > 3 else "",
                          "completion": clean_text(row[4]) if len(row) > 4 else "",
                          "progress": clean_text(row[5]) if len(row) > 5 else ""})
    if not items:
        return {"form": "trust_fund" if suffix == "6a" else "lgsf",
                "period": period_key(year, quarter), "year": year,
                "quarter": quarter, "status": "nil", "items": []}
    return {"form": "trust_fund" if suffix == "6a" else "lgsf",
            "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "items": items}


def parse_dfu_sheet(ws):
    rows = list(sheet_rows(ws, 12))
    year, quarter = parse_year_quarter(ws, rows)
    header_idx = next((i for i, r in enumerate(rows)
                       if clean_text(r[0]).lower().startswith("program or")), None)
    projects = []
    if header_idx is not None:
        for row in rows[header_idx + 2:]:
            a = clean_text(row[0]) if len(row) > 0 else ""
            if CERTIFY_RE.search(a) or a.lower().startswith("we hereby"):
                break
            if not a or a.upper() in ("N/A", "NONE"):
                continue
            if a.lower() in ("social development", "econimic development", "economic development"):
                continue
            projects.append({"project": a,
                             "location": clean_text(row[1]) if len(row) > 1 else "",
                             "cost": to_float(row[2]) if len(row) > 2 else None,
                             "started": clean_text(row[3]) if len(row) > 3 else "",
                             "completion": clean_text(row[4]) if len(row) > 4 else "",
                             "status": clean_text(row[5]) if len(row) > 5 else "",
                             "pct": to_float(row[6]) if len(row) > 6 else None,
                             "incurred": to_float(row[7]) if len(row) > 7 else None})
    return {"form": "dev_fund", "period": period_key(year, quarter), "year": year,
            "quarter": quarter, "sheet": ws.title, "projects": projects}


def parse_bids_sheet(ws, kind):
    rows = list(sheet_rows(ws, 12))
    year, quarter = parse_year_quarter(ws, rows)
    header_idx = next((i for i, r in enumerate(rows)
                       if "reference" in clean_text(r[1]).lower() or "reference" in clean_text(r[0]).lower()), None)
    bids = []
    if header_idx is not None:
        for row in rows[header_idx + 2:]:
            cells = [clean_text(c) for c in row]
            if CERTIFY_RE.search(" ".join(cells)) or cells[0].lower().startswith("we hereby"):
                break
            if not cells[1] or cells[1].upper() in ("N/A", "NONE"):
                continue
            bids.append({"ref": cells[1],
                         "project": clean_text(row[2]) if len(row) > 2 else "",
                         "abc": to_float(row[3]) if len(row) > 3 else None,
                         "location": clean_text(row[4]) if len(row) > 4 else "",
                         "bidder": clean_text(row[5]) if len(row) > 5 else "",
                         "bid_amount": to_float(row[7]) if len(row) > 7 else None,
                         "award_date": clean_text(row[8]) if len(row) > 8 else ""})
    return {"kind": kind, "bids": bids,
            "status": "nil" if not bids else "reported"}


def parse_bids_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    year = quarter = None
    parts = {}
    for ws in wb.worksheets:
        if ws.title == "FDPP LICENSE":
            continue
        low = ws.title.lower()
        kind = "civil_works" if "10a" in low else "goods" if "10b" in low else "consulting" if "10c" in low else None
        if not kind:
            continue
        parsed = parse_bids_sheet(ws, kind)
        parts[kind] = parsed["bids"]
        rows = list(sheet_rows(ws, 6))
        y, q = parse_year_quarter(ws, rows)
        year, quarter = y or year, q or quarter
    if not parts:
        return None
    return {"form": "bids", "period": period_key(year, quarter), "year": year,
            "quarter": quarter,
            "civil_works": parts.get("civil_works", []),
            "goods": parts.get("goods", []),
            "consulting": parts.get("consulting", [])}


def parse_manpower(ws):
    rows = list(sheet_rows(ws, 8))
    year, quarter = parse_year_quarter(ws, rows)
    items, total = [], None
    for row in rows:
        a = clean_text(row[0]) if len(row) > 0 else ""
        b = clean_text(row[1]) if len(row) > 1 else ""
        label = b or a
        low = label.lower()
        if low.startswith(("permanent", "contractual", "job order", "casual", "grand total")):
            count = to_float(row[2]) if len(row) > 2 else None
            amount = to_float(row[4]) if len(row) > 4 else None
            if low.startswith("grand total"):
                total = {"count": count, "amount": amount}
            else:
                items.append({"class": b, "count": count, "amount": amount})
    if not items:
        return None
    return {"form": "manpower", "period": period_key(year, quarter) if year else "undated",
            "year": year, "quarter": quarter, "items": items, "total": total,
            "note": None if year else "Snapshot date not stated in filing"}


def parse_sipb_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    loans = []
    for ws in wb.worksheets:
        items, report_date = {}, ""
        for row in sheet_rows(ws, 4):
            a = clean_text(row[0]) if len(row) > 0 else ""
            b = clean_text(row[1]) if len(row) > 1 else ""
            c = clean_text(row[2]) if len(row) > 2 else ""
            if a.isdigit() and b:
                items[b] = c
                if b.lower().startswith("date of report"):
                    report_date = c
        if items:
            loans.append({"loan": ws.title.strip(), "report_date": report_date, "items": items})
    return loans or None


def _detect_budget_year(ws, rows):
    for row in rows[:12]:
        joined = " ".join(clean_text(c) for c in row)
        match = re.search(r"[Bb][Uu][Dd][Gg][Ee][Tt]\s+[Yy][Ee][Aa][Rr]\s+(20\d{2})", joined)
        if match:
            return int(match.group(1))
        match = re.search(r"[Ff][Ii][Ss][Cc][Aa][Ll]\s+[Yy][Ee][Aa][Rr]\s+(20\d{2})", joined)
        if match:
            return int(match.group(1))
    match = re.search(r"(20\d{2})", ws.title)
    return int(match.group(1)) if match else None


def _parse_lbp2(ws, rows, year):
    offices, current = [], None
    for row in sheet_rows(ws, 13):
        a = clean_text(row[0]) if len(row) > 0 else ""
        if a.lower().startswith("office:"):
            current = {"office": re.sub(r"^office:\s*", "", a, flags=re.I),
                       "ps": 0.0, "mooe": 0.0, "co": 0.0, "proposed": 0.0, "lines": 0}
            offices.append(current)
            continue
        if current is None:
            continue
        code = clean_text(row[1]) if len(row) > 1 else ""
        if not CODE_RE.match(code):
            continue
        prop = to_float(row[11]) if len(row) > 11 else None
        if prop is None:
            continue
        current["lines"] += 1
        current["proposed"] += prop
        if code.startswith("5-01"):
            current["ps"] += prop
        elif code.startswith("5-02"):
            current["mooe"] += prop
        else:
            current["co"] += prop
    offices = [o for o in offices if o["lines"] > 0]
    if not offices:
        return None
    for o in offices:
        for k in ("ps", "mooe", "co", "proposed"):
            o[k] = round(o[k], 2)
    return {"form": "budget", "period": str(year) if year else "undated", "year": year,
            "quarter": None, "sheet": ws.title.strip(), "offices": offices,
            "totals": {k: round(sum(o[k] for o in offices), 2) for k in ("ps", "mooe", "co", "proposed")}}


def parse_spp_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    offices, summary = [], []
    year = None
    for ws in wb.worksheets:
        title = ws.title
        if title == "FDPP LICENSE" or title == "Sheet1":
            continue
        rows = list(sheet_rows(ws, 14))
        y, _ = parse_year_quarter(ws, rows)
        year = y or year
        if "14b" in title or "summary" in title.lower():
            for row in rows:
                a = clean_text(row[0]) if len(row) > 0 else ""
                if a.upper().startswith("OFFICE OF"):
                    total = to_float(row[4]) if len(row) > 4 else None
                    summary.append({"office": a, "total": total})
            continue
        office = ""
        for row in rows[:8]:
            cells = [clean_text(c) for c in row]
            for i, c in enumerate(cells):
                if c.lower() == "office:" and i + 1 < len(cells):
                    office = cells[i + 1]
        header_idx = next((i for i, r in enumerate(rows)
                           if "code (pap)" in clean_text(r[0]).lower()), None)
        items = []
        if header_idx is not None:
            for row in rows[header_idx + 2:]:
                cells = [clean_text(c) for c in row]
                if CERTIFY_RE.search(" ".join(cells)):
                    break
                if not cells[1] or cells[1].upper() in ("N/A", "NONE"):
                    continue
                amount = next((to_float(c) for c in reversed(cells) if to_float(c) is not None), None)
                items.append({"project": cells[1], "end_user": cells[2] if len(cells) > 2 else "",
                              "mode": cells[4] if len(cells) > 4 else "", "amount": amount})
        if office:
            offices.append({"office": office, "items": items})
    if not offices and not summary:
        return None
    return {"form": "spp", "period": str(year) if year else "undated", "year": year,
            "quarter": None, "offices": offices, "summary": summary}


# ----------------------------------------------------------------------------
# Driver
# ----------------------------------------------------------------------------

def sheet_signature(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return tuple(ws.title for ws in wb.worksheets if ws.title != "FDPP LICENSE")


def classify(path):
    sig = " | ".join(sheet_signature(path)).lower()
    if "sre_report" in sig and "soe_report" not in sig:
        return "sre"
    if "soe_report" in sig:
        return "sre_soe"
    if "sef utilization" in sig:
        return "sef"
    if "ldrrmfu" in sig:
        return "ldrrmf"
    if "form 9- scf" in sig:
        return "scf"
    if "form 12 - uca" in sig:
        return "uca"
    if "form 6a" in sig:
        return "trust"
    if "dfu" in sig and ("q3_2025" in sig or "form 7" in sig):
        return "dfu"
    if "form 10a" in sig:
        return "bids"
    if "lbp" in sig or "2026 final" in sig or "comparative" in sig:
        return "lbp"
    if "term loan" in sig:
        return "sipb"
    if "manpower" in sig or ("sheet1" in sig and "form 13" in (openpyxl.load_workbook(path, read_only=True, data_only=True)["Sheet1"]["A1"].value or "")):
        return "manpower"
    if "supplemental procurement" in sig or "14b" in sig:
        return "spp"
    if "sheet1" in sig:
        return "manpower"
    return "unknown:" + sig[:80]


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    if not CACHE_DIR.exists():
        sys.exit(f"missing {CACHE_DIR}")
    files = sorted(CACHE_DIR.rglob("*.xlsx"))
    print(f"found {len(files)} workbooks")

    seen_hashes = set()
    groups = {"sre": [], "sef": [], "ldrrmf": [], "cash_flows": [], "cash_advances": [],
              "trust_fund": [], "lgsf": [], "dev_fund": [], "bids": [], "manpower": [],
              "indebtedness": [], "budget": [], "budget_book": [], "spp": []}
    skipped, errors = [], []

    for path in files:
        digest = file_hash(path)
        if digest in seen_hashes:
            skipped.append(f"{path}: byte-duplicate")
            continue
        seen_hashes.add(digest)
        kind = classify(path)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if kind == "sre" or kind == "sre_soe":
                for ws in wb.worksheets:
                    if ws.title in ("sre_report", "soe_report"):
                        rec = parse_sre(ws)
                        if rec:
                            rec["sheet"] = ws.title
                            if ws.title == "soe_report":
                                rec["form"] = "soe"
                            rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                            groups["sre"].append(rec)
            elif kind == "sef":
                for ws in wb.worksheets:
                    if ws.title != "FDPP LICENSE":
                        rec = parse_sef(ws)
                        break
                else:
                    rec = None
                if rec:
                    rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                    groups["sef"].append(rec)
            elif kind == "ldrrmf":
                for ws in wb.worksheets:
                    if "LDRRMFU" in ws.title:
                        rec = parse_ldrrmf(ws)
                        if rec:
                            rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                            groups["ldrrmf"].append(rec)
            elif kind == "scf":
                for ws in wb.worksheets:
                    if "SCF" in ws.title:
                        rec = parse_scf(ws)
                        if rec:
                            rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                            groups["cash_flows"].append(rec)
            elif kind == "uca":
                for ws in wb.worksheets:
                    if "UCA" in ws.title:
                        rec = parse_uca(ws)
                        if rec:
                            rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                            groups["cash_advances"].append(rec)
            elif kind == "trust":
                for ws in wb.worksheets:
                    if "6a" in ws.title:
                        rec = parse_trust(ws, "6a")
                    elif "6b" in ws.title:
                        rec = parse_trust(ws, "6b")
                    else:
                        continue
                    rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                    groups[rec["form"]].append(rec)
            elif kind == "dfu":
                for ws in wb.worksheets:
                    if ws.title == "FDPP LICENSE":
                        continue
                    rec = parse_dfu_sheet(ws)
                    rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}#{ws.title}"
                    groups["dev_fund"].append(rec)
            elif kind == "bids":
                rec = parse_bids_file(path)
                if rec:
                    rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                    groups["bids"].append(rec)
            elif kind == "manpower":
                for ws in wb.worksheets:
                    if ws.title == "FDPP LICENSE":
                        continue
                    rec = parse_manpower(ws)
                    if rec:
                        rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                        groups["manpower"].append(rec)
            elif kind == "sipb":
                loans = parse_sipb_file(path)
                if loans:
                    for loan in loans:
                        loan["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                    groups["indebtedness"].extend(loans)
            elif kind == "lbp":
                try:
                    default_year = int(path.parent.name)
                except ValueError:
                    default_year = None
                for ws in wb.worksheets:
                    rec = _parse_lbp_sheet(ws, default_year)
                    if rec:
                        rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}#{ws.title}"
                        groups[rec["form"]].append(rec)
            elif kind == "spp":
                rec = parse_spp_file(path)
                if rec:
                    rec["source_file"] = f"datasets/fdp/{path.parent.name}/{path.name}"
                    groups["spp"].append(rec)
            else:
                skipped.append(f"{path}: unclassified ({kind})")
        except (OSError, ValueError, KeyError, AttributeError, zipfile.BadZipFile) as exc:
            errors.append(f"{path}: {type(exc).__name__}: {exc}")

    # dedup: same (form, period) from different files -> keep most data rows.
    # Undated snapshots and explicit nil filings are distinguishable only by
    # source file, so they are never collapsed.
    for key in ("sre", "sef", "ldrrmf", "cash_flows", "cash_advances", "trust_fund",
                "lgsf", "bids", "manpower", "spp", "budget", "budget_book", "dev_fund"):
        best, keep_all = {}, []
        for rec in groups[key]:
            if rec.get("period") in (None, "undated") or rec.get("status") == "nil":
                keep_all.append(rec)
                continue
            pk = (rec.get("period"), rec.get("sheet", ""))
            rows = len(rec.get("rows", []) or rec.get("items", []) or rec.get("projects", [])
                       or rec.get("civil_works", []) or rec.get("goods", []) or rec.get("offices", [])
                       or rec.get("debtors", []) or [])
            if pk not in best or rows > best[pk][0]:
                best[pk] = (rows, rec)
        groups[key] = sorted([rec for _, rec in best.values()] + keep_all,
                             key=lambda r: (r.get("period", ""), r.get("source_file", "")))
    # dev_fund: same quarter may arrive on differently-named sheets across
    # files — merge projects by period instead of picking one sheet.
    merged = {}
    for rec in groups["dev_fund"]:
        period = rec.get("period", "undated")
        slot = merged.setdefault(period, {"form": "dev_fund", "period": period,
                                          "year": rec.get("year"), "quarter": rec.get("quarter"),
                                          "projects": [], "source_file": []})
        seen = {(p.get("project"), p.get("location"), p.get("cost")) for p in slot["projects"]}
        for proj in rec.get("projects", []):
            key = (proj.get("project"), proj.get("location"), proj.get("cost"))
            if key not in seen:
                seen.add(key)
                slot["projects"].append(proj)
        src = rec.get("source_file", "")
        if src and src not in slot["source_file"]:
            slot["source_file"].append(src)
    groups["dev_fund"] = [merged[p] for p in sorted(merged)]

    data = {"meta": {"municipality": "Mapandan, Pangasinan",
                     "source": "DILG Full Disclosure Policy Portal",
                     "retrieved": "2026-09-18",
                     "years": sorted({r.get("year") for g in groups.values() for r in g if r.get("year")})},
            **groups}
    write_json(OUTPUT, data)
    print(f"wrote {OUTPUT}")
    for key, recs in groups.items():
        print(f"  {key}: {len(recs)} records")
    for line in skipped:
        print("  SKIP:", line)
    for line in errors:
        print("  ERROR:", line)


def _parse_lbp_sheet(ws, default_year=None):
    rows = list(sheet_rows(ws, 14))
    year = _detect_budget_year(ws, rows) or default_year
    has_account_code = any(len(r) > 1 and "account code" in clean_text(r[1]).lower() for r in rows[:14])
    if has_account_code:
        return _parse_lbp2(ws, rows, year)
    # Generic budget-book sheet (LBP 1/5/6/7, COMPARATIVE, per-dept matrix):
    # capture labeled rows with their trailing amounts.
    captured = []
    for row in rows:
        label = clean_text(row[0]) if len(row) > 0 else ""
        if not label or len(label) > 120:
            continue
        low = label.lower()
        if low.startswith(("office:", "object of", "particip", "region", "province", "city",
                            "municipality", "mapandan", "general fund", "lbp form", "budget of",
                            "programmed", "statement of", "comparative", "appropriations for")):
            continue
        if CERTIFY_RE.search(label):
            continue
        vals = [to_float(c) for c in row[1:]]
        vals = [v for v in vals if v is not None]
        if vals:
            captured.append({"label": label, "values": vals})
    if not captured:
        return None
    return {"form": "budget_book", "sheet": ws.title.strip(),
            "period": str(year) if year else "undated", "year": year,
            "quarter": None, "rows": captured}


if __name__ == "__main__":
    main()
