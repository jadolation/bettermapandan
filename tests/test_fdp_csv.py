"""Tests for the FDP xlsx -> csv dump script (local analysis output)."""
import csv
import sys
from datetime import date, datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.data.fdp_to_csv import cell_to_text, convert_workbook, dump_sheet, sanitize


def test_sanitize_removes_unsafe_chars():
    assert sanitize("REGION I - ILOCOS REGION PANGASINAN MAPANDAN (42)") == (
        "REGION_I_-_ILOCOS_REGION_PANGASINAN_MAPANDAN_42"
    )


def test_cell_to_text_values():
    assert cell_to_text(None) == ""
    assert cell_to_text(True) == "TRUE"
    assert cell_to_text(42) == "42"
    assert cell_to_text(1.5) == "1.5"
    assert cell_to_text(datetime(2025, 5, 5, 10, 30, tzinfo=timezone.utc)) == (
        "2025-05-05 10:30:00+00:00"
    )
    assert cell_to_text("  padded\ttext  ") == "padded text"


def _make_workbook(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SRE Report"
    ws["A1"] = "Particulars"
    ws["B2"] = 1234.5
    ws["C3"] = date(2026, 5, 5)
    ws["D4"] = None
    lic = wb.create_sheet("FDPP LICENSE")
    lic["A1"] = "boilerplate"
    wb.save(path)


def test_convert_workbook_skips_license_and_roundtrips(tmp_path):
    xlsx = tmp_path / "sample.xlsx"
    _make_workbook(xlsx)
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    written, skipped = convert_workbook(xlsx, out_dir)

    assert (written, skipped) == (1, 1)
    csv_path = out_dir / "sample__SRE_Report.csv"
    assert csv_path.is_file()
    with csv_path.open(encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows[0][0] == "Particulars"
    assert rows[1][1] == "1234.5"
    assert rows[2][2] == "2026-05-05 00:00:00"


def test_convert_is_up_to_date_skip(tmp_path):
    xlsx = tmp_path / "sample.xlsx"
    _make_workbook(xlsx)
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    assert convert_workbook(xlsx, out_dir) == (1, 1)
    assert convert_workbook(xlsx, out_dir) == (1, 1)  # Fresh; counted, not rewritten.


def test_dump_sheet_trims_empty_edges(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["C2"] = "y"
    csv_path = tmp_path / "trim.csv"
    assert dump_sheet(ws, csv_path) == 2
    with csv_path.open(encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows == [["x", "", ""], ["", "", "y"]]
